"""Immutable per-device production reports and self-contained evidence archives."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import base64
import csv
import hashlib
from io import BytesIO
import json
from pathlib import Path
import platform
import re
from typing import Any, Mapping, Optional, Sequence
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from satellite_debug_tool import __version__

from .models import AttemptStatus, BatchStatus
from .report_analysis import build_report_analysis, render_report_charts
from .result_store import ProductionResultStore, ResultStoreError


_TERMINAL_ATTEMPT_STATUSES = {
    AttemptStatus.PASS.value,
    AttemptStatus.FAIL.value,
    AttemptStatus.INCOMPLETE.value,
    AttemptStatus.SKIPPED.value,
    AttemptStatus.ABORTED.value,
}
_FORMAL_VERDICTS = {"PASS", "FAIL"}
_INVALID_PATH_CHARACTERS = re.compile(r"[<>:\"/\\|?*\x00-\x1f]")


@dataclass(frozen=True)
class ReportBranding:
    company_name: str = ""
    logo_path: str = ""
    header: str = ""
    footer: str = ""
    tester_role: str = ""
    reviewer_role: str = ""
    logo_base64: str = ""
    logo_filename: str = ""

    @classmethod
    def from_mapping(cls, value: Optional[Mapping[str, Any]]) -> "ReportBranding":
        payload = value or {}
        return cls(
            company_name=str(payload.get("company_name", "") or "").strip(),
            logo_path=str(payload.get("logo_path", "") or "").strip(),
            header=str(payload.get("header", "") or "").strip(),
            footer=str(payload.get("footer", "") or "").strip(),
            tester_role=str(payload.get("tester_role", "") or "").strip(),
            reviewer_role=str(payload.get("reviewer_role", "") or "").strip(),
            logo_base64=str(payload.get("logo_base64", "") or "").strip(),
            logo_filename=str(payload.get("logo_filename", "") or "").strip(),
        )

    def to_mapping(self) -> dict[str, str]:
        return {
            "company_name": self.company_name,
            "logo_path": self.logo_path,
            "header": self.header,
            "footer": self.footer,
            "tester_role": self.tester_role,
            "reviewer_role": self.reviewer_role,
            "logo_base64": self.logo_base64,
            "logo_filename": self.logo_filename,
        }

    @property
    def missing_fields(self) -> tuple[str, ...]:
        fields = {
            "公司名称": self.company_name,
            "Logo": self.logo_path or self.logo_base64,
            "测试员角色": self.tester_role,
            "审核员角色": self.reviewer_role,
        }
        return tuple(name for name, value in fields.items() if not value)


@dataclass(frozen=True)
class ReportArtifact:
    batch_id: str
    serial_number: str
    report_id: str
    version: int
    status: str
    automatic_verdict: str
    final_verdict: str
    directory: Path
    document_path: Optional[Path]
    evidence_path: Path
    manifest_path: Path


def sanitize_path_component(value: object, *, fallback: str = "unknown") -> str:
    text = _INVALID_PATH_CHARACTERS.sub("_", str(value or "").strip())
    text = re.sub(r"\s+", "_", text).strip(" ._")
    if text in {"", ".", ".."}:
        return fallback
    return text[:128]


def automatic_device_verdict(attempts: Sequence[Mapping[str, Any]]) -> str:
    statuses = {str(item.get("status", "")) for item in attempts}
    if not statuses or any(status not in _TERMINAL_ATTEMPT_STATUSES for status in statuses):
        return "INCOMPLETE"
    if AttemptStatus.ABORTED.value in statuses:
        return "ABORTED"
    if AttemptStatus.INCOMPLETE.value in statuses:
        return "INCOMPLETE"
    if AttemptStatus.FAIL.value in statuses:
        return "FAIL"
    if AttemptStatus.PASS.value in statuses:
        return "PASS"
    return "INCOMPLETE"


class ProductionReportService:
    """Create V1/V2 reports exclusively from a frozen device snapshot."""

    def __init__(
        self,
        store: ProductionResultStore,
        *,
        batch_output_dir: str | Path,
        report_root: str | Path,
        branding: ReportBranding = ReportBranding(),
    ) -> None:
        self.store = store
        self.batch_output_dir = Path(batch_output_dir).expanduser().resolve()
        self.report_root = Path(report_root).expanduser().resolve()
        self.branding = branding

    def archive_batch(self, batch_id: str) -> tuple[ReportArtifact, ...]:
        batch = self.store.get_batch(batch_id)
        if batch["status"] not in {
            BatchStatus.COMPLETED.value,
            BatchStatus.INCOMPLETE.value,
            BatchStatus.ABORTED.value,
        }:
            raise ResultStoreError("reports require a terminal batch")
        artifacts = []
        for device in self.store.list_devices(batch_id):
            if device["status"] != "participant":
                continue
            artifacts.append(self.archive_device(batch_id, device["serial_number"]))
        return tuple(artifacts)

    def archive_device(self, batch_id: str, serial_number: str) -> ReportArtifact:
        try:
            frozen = self.store.get_report_snapshot(batch_id, serial_number)
            snapshot = frozen["snapshot_json"]
        except ResultStoreError:
            snapshot = self._build_snapshot(batch_id, serial_number)
            frozen = self.store.freeze_report_snapshot(
                batch_id,
                serial_number,
                snapshot,
                automatic_verdict=snapshot["automatic_verdict"],
                report_id=snapshot["report_id"],
            )
        verdict = str(frozen["automatic_verdict"])
        directory = self._artifact_directory(snapshot)
        directory.mkdir(parents=True, exist_ok=True)
        evidence_path, evidence_entries = self._write_evidence_zip(snapshot, directory)
        document_path: Optional[Path] = None
        status = "generated_v1" if verdict in _FORMAL_VERDICTS else "not_generated"
        try:
            if verdict in _FORMAL_VERDICTS:
                document_path = directory / f"{snapshot['report_id']}_V01.docx"
                if document_path.exists():
                    records = self.store.list_report_artifacts(batch_id, serial_number)
                    failed = any(
                        item["version"] == 1 and item["status"] == "generation_failed"
                        for item in records
                    )
                    if not failed:
                        raise ResultStoreError(
                            f"report artifact already exists: {document_path.name}"
                        )
                    document_path.unlink()
                self._write_docx(snapshot, document_path, version=1, review=None)
            manifest_path = self._write_manifest(
                snapshot,
                directory,
                evidence_path,
                evidence_entries,
            )
            if verdict in _FORMAL_VERDICTS:
                self.store.record_report_artifact(
                    batch_id,
                    serial_number,
                    version=1,
                    status="generated_v1",
                    final_verdict=verdict,
                    document_path=str(document_path),
                    document_sha256=_sha256_file(document_path),
                    evidence_path=str(evidence_path),
                    evidence_sha256=_sha256_file(evidence_path),
                    manifest_path=str(manifest_path),
                )
            else:
                self.store.record_report_artifact(
                    batch_id,
                    serial_number,
                    version=1,
                    status="not_generated",
                    final_verdict=verdict,
                    evidence_path=str(evidence_path),
                    evidence_sha256=_sha256_file(evidence_path),
                    manifest_path=str(manifest_path),
                )
        except Exception as exc:
            if verdict in _FORMAL_VERDICTS:
                self.store.record_report_artifact(
                    batch_id,
                    serial_number,
                    version=1,
                    status="generation_failed",
                    final_verdict=verdict,
                    evidence_path=str(evidence_path),
                    evidence_sha256=_sha256_file(evidence_path),
                    error=str(exc),
                )
            raise
        return ReportArtifact(
            batch_id=batch_id,
            serial_number=serial_number,
            report_id=snapshot["report_id"],
            version=1,
            status=status,
            automatic_verdict=verdict,
            final_verdict=verdict,
            directory=directory,
            document_path=document_path,
            evidence_path=evidence_path,
            manifest_path=manifest_path,
        )

    def review(
        self,
        batch_id: str,
        serial_number: str,
        *,
        reviewer: str,
        final_verdict: str,
        reason: str,
    ) -> ReportArtifact:
        try:
            review = self.store.get_report_review(batch_id, serial_number)
        except ResultStoreError:
            review = self.store.create_report_review(
                batch_id,
                serial_number,
                reviewer=reviewer,
                final_verdict=final_verdict,
                reason=reason,
            )
        else:
            if (
                review["reviewer"] != str(reviewer).strip()
                or review["final_verdict"] != str(final_verdict).upper()
                or review["reason"] != str(reason).strip()
            ):
                raise ResultStoreError("report review already exists with different content")
        frozen = self.store.get_report_snapshot(batch_id, serial_number)
        snapshot = frozen["snapshot_json"]
        directory = self._artifact_directory(snapshot)
        evidence_path = directory / "evidence.zip"
        if not evidence_path.is_file():
            raise ResultStoreError("review requires the original evidence.zip")
        document_path = directory / f"{snapshot['report_id']}_V02.docx"
        if document_path.exists():
            records = self.store.list_report_artifacts(batch_id, serial_number)
            failed = any(
                item["version"] == 2 and item["status"] == "generation_failed"
                for item in records
            )
            if not failed:
                raise ResultStoreError(f"report artifact already exists: {document_path.name}")
            document_path.unlink()
        try:
            self._write_docx(snapshot, document_path, version=2, review=review)
            manifest_path = self._write_manifest(
                snapshot,
                directory,
                evidence_path,
                _zip_entry_hashes(evidence_path),
                review=review,
            )
            self.store.record_report_artifact(
                batch_id,
                serial_number,
                version=2,
                status="reviewed_v2",
                final_verdict=review["final_verdict"],
                document_path=str(document_path),
                document_sha256=_sha256_file(document_path),
                evidence_path=str(evidence_path),
                evidence_sha256=_sha256_file(evidence_path),
                manifest_path=str(manifest_path),
            )
        except Exception as exc:
            self.store.record_report_artifact(
                batch_id,
                serial_number,
                version=2,
                status="generation_failed",
                final_verdict=review["final_verdict"],
                evidence_path=str(evidence_path),
                evidence_sha256=_sha256_file(evidence_path),
                error=str(exc),
            )
            raise
        return ReportArtifact(
            batch_id=batch_id,
            serial_number=serial_number,
            report_id=snapshot["report_id"],
            version=2,
            status="reviewed_v2",
            automatic_verdict=snapshot["automatic_verdict"],
            final_verdict=review["final_verdict"],
            directory=directory,
            document_path=document_path,
            evidence_path=evidence_path,
            manifest_path=manifest_path,
        )

    def _build_snapshot(self, batch_id: str, serial_number: str) -> dict[str, Any]:
        batch = self.store.get_batch(batch_id)
        device = self.store.get_device(batch_id, serial_number)
        attempts = [
            item
            for item in self.store.list_attempts(batch_id)
            if item["serial_number"] == serial_number
        ]
        for attempt in attempts:
            attempt["metrics"] = self.store.list_metrics(int(attempt["attempt_id"]))
        verdict = automatic_device_verdict(attempts)
        finished = _parse_datetime(batch.get("finished_utc") or batch["created_utc"])
        date_text = finished.astimezone().strftime("%Y%m%d")
        model = str(device.get("hardware_type") or batch["recipe_json"].get("product") or "unknown").upper()
        safe_model = sanitize_path_component(model)
        safe_serial = sanitize_path_component(serial_number)
        safe_batch = sanitize_path_component(batch_id)
        report_id = f"RPT-{date_text}-{safe_model}-{safe_serial}-{safe_batch}"
        configuration = _read_json_file(self.batch_output_dir / "production_configuration.json")
        branding = self.branding.to_mapping()
        configured_branding = configuration.get("station_profile", {}).get("report_branding")
        if isinstance(configured_branding, Mapping):
            branding.update(
                {
                    str(key): str(value)
                    for key, value in configured_branding.items()
                    if value
                }
            )
        events = self.store.list_events(batch_id)
        sdb_files = self._device_sdb_files_for_serial(serial_number)
        analysis = build_report_analysis(
            sdb_files,
            attempts,
            events,
            self.batch_output_dir,
        )
        local_now = datetime.now().astimezone()
        return {
            "schema": "satellite-debug-tool/device-report-snapshot",
            "schema_version": 2,
            "report_id": report_id,
            "automatic_verdict": verdict,
            "frozen_utc": datetime.now(timezone.utc).isoformat(timespec="milliseconds"),
            "batch": batch,
            "device": device,
            "attempts": attempts,
            "fixture_actions": self.store.list_fixture_actions(batch_id),
            "events": events,
            "analysis": analysis,
            "configuration": configuration,
            "branding": branding,
            "application": {
                "version": __version__,
                "python": platform.python_version(),
                "platform": platform.platform(),
                "hostname": platform.node(),
                "timezone": str(local_now.tzinfo or "未配置"),
                "utc_offset": local_now.strftime("%z") or "未配置",
            },
        }

    def _artifact_directory(self, snapshot: Mapping[str, Any]) -> Path:
        batch = snapshot["batch"]
        device = snapshot["device"]
        finished = _parse_datetime(batch.get("finished_utc") or batch["created_utc"])
        date_dir = finished.astimezone().strftime("%Y-%m-%d")
        model = str(device.get("hardware_type") or batch["recipe_json"].get("product") or "unknown").upper()
        return (
            self.report_root
            / date_dir
            / sanitize_path_component(model)
            / sanitize_path_component(device["serial_number"])
            / sanitize_path_component(batch["batch_id"])
        )

    def _write_evidence_zip(
        self, snapshot: Mapping[str, Any], directory: Path
    ) -> tuple[Path, dict[str, str]]:
        entries = self._evidence_entries(snapshot)
        target = directory / "evidence.zip"
        temporary = target.with_name(f".{target.name}.tmp")
        hashes: dict[str, str] = {}
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, payload in entries.items():
                data = payload if isinstance(payload, bytes) else payload.encode("utf-8")
                archive.writestr(name, data)
                hashes[name] = hashlib.sha256(data).hexdigest()
            for index, source in enumerate(self._device_sdb_files(snapshot), start=1):
                name = "device/evidence.sdb" if index == 1 else f"device/evidence-{index:02d}.sdb"
                archive.write(source, name)
                hashes[name] = _sha256_file(source)
        temporary.replace(target)
        return target, hashes

    def _evidence_entries(self, snapshot: Mapping[str, Any]) -> dict[str, bytes | str]:
        device = snapshot["device"]
        attempts = snapshot["attempts"]
        metrics = [metric for attempt in attempts for metric in attempt.get("metrics", [])]
        configuration = snapshot.get("configuration", {})
        event_lines = "".join(_json_line(item) for item in snapshot["events"])
        fixture_actions = snapshot.get("fixture_actions", [])
        power_records = _read_text_file(self.batch_output_dir / "power_supply_records.jsonl")
        power_summary = _read_json_file(self.batch_output_dir / "power_supply_summary.json")
        return {
            "device/identity.json": _pretty_json(device),
            "device/parameters.json": _pretty_json(
                device.get("metadata_json", {}).get("parameters", {"status": "not_recorded"})
            ),
            "device/metrics.json": _pretty_json(metrics),
            "device/events.jsonl": event_lines,
            "fixture/power_config.json": _pretty_json(configuration.get("power_profile", {})),
            "fixture/power_identity.json": _pretty_json(power_summary.get("identity", {"status": "unavailable"})),
            "fixture/power_scpi.jsonl": power_records or _json_line({"status": "unavailable"}),
            "fixture/motion_commands.jsonl": _fixture_lines(fixture_actions, "motion"),
            "fixture/ms6222_raw.jsonl": _first_text(
                self.batch_output_dir,
                ("ms6222_raw.jsonl", "raw_frames.jsonl"),
            ),
            "fixture/ms6222_parsed.jsonl": _first_text_or_csv_jsonl(
                self.batch_output_dir,
                ("ms6222_parsed.jsonl", "parsed_samples.jsonl", "comparisons.jsonl"),
                ("ms6222_parsed.csv", "parsed_frames.csv", "attitude_comparison.csv"),
            ),
            "fixture/attitude_comparison.csv": _first_text(
                self.batch_output_dir,
                ("attitude_comparison.csv",),
            ),
            "fixture/calibration.json": _pretty_json(
                _first_json(self.batch_output_dir, ("calibration.json", "ms6222_calibration.json"))
            ),
            "test/recipe.json": _pretty_json(snapshot["batch"]["recipe_json"]),
            "test/attempts.json": _pretty_json(attempts),
            "test/timeline.json": _pretty_json(snapshot["events"]),
            "test/software.json": _pretty_json(
                {
                    **snapshot["application"],
                    "batch_id": snapshot["batch"]["batch_id"],
                    "report_id": snapshot["report_id"],
                    "snapshot_frozen_utc": snapshot["frozen_utc"],
                }
            ),
            "analysis/analysis.json": _pretty_json(snapshot.get("analysis", {})),
            **render_report_charts(snapshot.get("analysis", {})),
        }

    def _device_sdb_files(self, snapshot: Mapping[str, Any]) -> tuple[Path, ...]:
        return self._device_sdb_files_for_serial(str(snapshot["device"]["serial_number"]))

    def _device_sdb_files_for_serial(self, serial: str) -> tuple[Path, ...]:
        candidates = self.batch_output_dir / "devices" / serial / "evidence"
        if not candidates.is_dir():
            return ()
        return tuple(sorted(path for path in candidates.glob("*.sdb") if path.is_file()))

    def _write_manifest(
        self,
        snapshot: Mapping[str, Any],
        directory: Path,
        evidence_path: Path,
        evidence_entries: Mapping[str, str],
        *,
        review: Optional[Mapping[str, Any]] = None,
    ) -> Path:
        reports = []
        for path in sorted(directory.glob(f"{snapshot['report_id']}_V*.docx")):
            reports.append({"name": path.name, "sha256": _sha256_file(path)})
        payload = {
            "schema": "satellite-debug-tool/device-report-manifest",
            "schema_version": 1,
            "generated_utc": datetime.now(timezone.utc).isoformat(timespec="milliseconds"),
            "application_version": __version__,
            "batch_id": snapshot["batch"]["batch_id"],
            "serial_number": snapshot["device"]["serial_number"],
            "report_id": snapshot["report_id"],
            "automatic_verdict": snapshot["automatic_verdict"],
            "final_verdict": (
                review["final_verdict"] if review is not None else snapshot["automatic_verdict"]
            ),
            "snapshot_sha256": hashlib.sha256(
                json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
            ).hexdigest(),
            "evidence": {
                "name": evidence_path.name,
                "sha256": _sha256_file(evidence_path),
                "entries": dict(sorted(evidence_entries.items())),
            },
            "reports": reports,
            "review": None if review is None else dict(review),
        }
        target = directory / "manifest.json"
        _atomic_write_text(target, _pretty_json(payload))
        return target

    def _write_docx(
        self,
        snapshot: Mapping[str, Any],
        destination: Path,
        *,
        version: int,
        review: Optional[Mapping[str, Any]],
    ) -> None:
        document = Document()
        _configure_document(document, snapshot.get("branding", {}))
        _add_cover(document, snapshot, version, review)
        document.add_page_break()
        _add_summary(document, snapshot, review)
        document.add_page_break()
        _add_configuration(document, snapshot)
        document.add_page_break()
        _add_curve_analysis(document, snapshot)
        _add_results(document, snapshot)
        document.add_page_break()
        _add_integrity(document, snapshot)
        document.add_page_break()
        _add_review(document, snapshot, version, review)
        document.add_page_break()
        _add_appendix(document, snapshot)
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(f".{destination.name}.tmp")
        document.save(temporary)
        temporary.replace(destination)


def _configure_document(document: Document, branding: Mapping[str, Any]) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Arial Unicode MS"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    header = section.header.paragraphs[0]
    header.text = str(branding.get("header") or branding.get("company_name") or "未配置")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(str(branding.get("footer") or "卫星通信相控阵设备自动化测试报告  |  "))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _add_cover(
    document: Document,
    snapshot: Mapping[str, Any],
    version: int,
    review: Optional[Mapping[str, Any]],
) -> None:
    branding = snapshot.get("branding", {})
    batch = snapshot["batch"]
    device = snapshot["device"]
    configuration = snapshot.get("configuration", {})
    product = configuration.get("product_template", {})
    verdict = review["final_verdict"] if review else snapshot["automatic_verdict"]

    report_number = document.add_paragraph()
    report_number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    report_number.paragraph_format.space_after = Pt(28)
    number_run = report_number.add_run(str(snapshot["report_id"]))
    number_run.font.size = Pt(8)
    number_run.font.color.rgb = RGBColor(89, 89, 89)
    _set_paragraph_bottom_border(report_number)

    logo_path = Path(str(branding.get("logo_path", "") or "")).expanduser()
    logo_source: Any = None
    if logo_path.is_file():
        logo_source = str(logo_path)
    elif branding.get("logo_base64"):
        try:
            logo_source = BytesIO(base64.b64decode(str(branding["logo_base64"]), validate=True))
        except (ValueError, TypeError):
            logo_source = None
    if logo_source is not None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run().add_picture(logo_source, width=Cm(2.6))
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("测试报告")
    run.bold = True
    run.font.size = Pt(26)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    subtitle_run = subtitle.add_run("卫星通信相控阵设备自动化测试")
    subtitle_run.font.size = Pt(11)

    product_name = product.get("display_name") or device.get("hardware_type") or "未配置"
    product_model = (
        device.get("hardware_type")
        or product.get("product")
        or batch["recipe_json"].get("product")
        or "未配置"
    )
    for label, value in (
        ("产品名称", product_name),
        ("产品型号", product_model),
        ("产品编号", device["serial_number"]),
        ("批次编号", batch["batch_id"]),
        ("报告版本", f"V{version:02d}"),
        ("测试结论", verdict),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        label_run = paragraph.add_run(f"{label}：  ")
        label_run.bold = True
        label_run.font.size = Pt(12)
        value_run = paragraph.add_run(str(value))
        value_run.bold = True
        value_run.font.size = Pt(12)

    document.add_paragraph().paragraph_format.space_after = Pt(6)
    writer = str(batch.get("operator") or "未配置")
    tester_role = str(branding.get("tester_role") or "").strip()
    if tester_role:
        writer = f"{writer}（{tester_role}）"
    reviewer = "未审核" if review is None else str(review["reviewer"])
    reviewer_role = str(branding.get("reviewer_role") or "").strip()
    if review is not None and reviewer_role:
        reviewer = f"{reviewer}（{reviewer_role}）"
    _add_signature_table(
        document,
        (("编写", writer), ("审核", reviewer), ("批准", "未配置")),
    )

    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.paragraph_format.space_before = Pt(26)
    date_run = date_paragraph.add_run(_cover_date(batch))
    date_run.bold = True
    date_run.font.size = Pt(12)


def _add_summary(document: Document, snapshot: Mapping[str, Any], review: Optional[Mapping[str, Any]]) -> None:
    document.add_heading("一、测试摘要", level=1)
    attempts = snapshot["attempts"]
    counts = {status: 0 for status in ("pass", "fail", "incomplete", "skipped", "aborted")}
    for attempt in attempts:
        counts.setdefault(attempt["status"], 0)
        counts[attempt["status"]] += 1
    final_verdict = review["final_verdict"] if review else snapshot["automatic_verdict"]
    _add_key_value_table(
        document,
        (
            ("自动结论", snapshot["automatic_verdict"]),
            ("最终结论", final_verdict),
            ("测试时段", _time_range(snapshot["batch"])),
            ("测试项总数", len(attempts)),
            ("通过 / 失败", f"{counts['pass']} / {counts['fail']}"),
            ("不完整 / 中止 / N/A", f"{counts['incomplete']} / {counts['aborted']} / {counts['skipped']}"),
        ),
    )
    branding = ReportBranding.from_mapping(snapshot.get("branding"))
    if branding.missing_fields:
        document.add_paragraph(
            "报告信息不完整：" + "、".join(branding.missing_fields),
            style="Intense Quote",
        )
    rows = [("测试项目", "次数", "自动结果", "摘要")]
    grouped: dict[str, list[Mapping[str, Any]]] = {}
    for attempt in attempts:
        grouped.setdefault(attempt["test_id"], []).append(attempt)
    for test_id, values in grouped.items():
        latest = max(values, key=lambda item: int(item["ordinal"]))
        rows.append(
            (
                test_id,
                str(len(values)),
                str(latest["status"]).upper(),
                _compact_json(latest.get("result_json", {})),
            )
        )
    _add_table(document, rows, widths=(5.0, 2.0, 3.0, 7.0))
    document.add_heading("过程数据总结", level=2)
    document.add_paragraph(
        str(snapshot.get("analysis", {}).get("summary") or "未生成过程曲线分析。")
    )


def _add_configuration(document: Document, snapshot: Mapping[str, Any]) -> None:
    document.add_heading("二、设备与测试配置", level=1)
    batch = snapshot["batch"]
    device = snapshot["device"]
    metadata = device.get("metadata_json", {})
    configuration = snapshot.get("configuration", {})
    power = configuration.get("power_profile", {})
    station = configuration.get("station_profile", {})
    combined = configuration.get("combined_supply", {})
    environment = configuration.get("test_environment", {})
    application = snapshot.get("application", {})
    _add_key_value_table(
        document,
        (
            ("设备型号", device.get("hardware_type") or "未配置"),
            ("设备 SN", device["serial_number"]),
            ("设备 UID", device.get("device_uid") or "未配置"),
            ("MAC", device.get("mac_address") or "未配置"),
            ("网络端点", f"{device.get('endpoint_ip', '')}:{device.get('endpoint_port', '')}"),
            ("主固件", metadata.get("main_firmware") or "未配置"),
            ("配方", batch["recipe_json"].get("recipe_id", "未配置")),
            ("配方 SHA-256", batch.get("recipe_sha256", "")),
            ("电源", f"{power.get('manufacturer', '未配置')} {power.get('model', '')}"),
            ("电源序列号", power.get("serial_number") or "未配置"),
            ("批次电压 / 限流", f"{combined.get('voltage_v', '未配置')} V / {combined.get('current_a', '未配置')} A"),
            ("摇摆台配置", station.get("motion_profile_id") or "未配置"),
            ("MS6222 配置", station.get("reference_profile_id") or "未配置"),
            ("工作站", station.get("display_name") or platform.node() or "未配置"),
        ),
    )
    document.add_heading("测试环境", level=2)
    power_endpoint = _network_endpoint(power.get("host"), power.get("port"))
    _add_key_value_table(
        document,
        (
            ("测试地点", environment.get("location") or "未配置"),
            ("测试工位", station.get("display_name") or "未配置"),
            ("工作站主机", application.get("hostname") or "未配置"),
            ("操作系统", application.get("platform") or "未配置"),
            ("测试软件", f"satellite_debug_tool {application.get('version', '未配置')}"),
            ("运行时", f"Python {application.get('python', '未配置')}"),
            (
                "测试时区",
                f"{application.get('timezone', '未配置')} (UTC{_format_utc_offset(application.get('utc_offset'))})",
            ),
            ("设备网络端点", _network_endpoint(device.get("endpoint_ip"), device.get("endpoint_port"))),
            ("电源连接", power_endpoint),
            ("摇摆台配置", station.get("motion_profile_id") or "未配置"),
            ("MS6222 配置", station.get("reference_profile_id") or "未配置"),
            ("环境温度", _environment_value(environment.get("temperature_c"), "°C")),
            ("相对湿度", _environment_value(environment.get("relative_humidity_percent"), "%RH")),
        ),
    )
    document.add_heading("参数设定与回读", level=2)
    parameters = metadata.get("parameters")
    document.add_paragraph(_pretty_json(parameters if parameters is not None else {"status": "not_recorded"}))


def _add_curve_analysis(document: Document, snapshot: Mapping[str, Any]) -> None:
    document.add_heading("三、过程曲线与阶段分析", level=1)
    analysis = snapshot.get("analysis", {})
    series = analysis.get("series", {})
    document.add_paragraph(
        str(analysis.get("summary") or "未提取到可分析的过程数据。")
    )
    charts = render_report_charts(analysis)
    for name, caption in (
        ("analysis/snr_timeline.png", "图 3-1  全过程 SNR 与锁定状态"),
        ("analysis/attitude_timeline.png", "图 3-2  全过程设备姿态"),
        ("analysis/convergence_timeline.png", "图 3-3  MS6222 姿态误差与收敛/发散趋势"),
    ):
        payload = charts.get(name)
        if not payload:
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(BytesIO(payload), width=Cm(16.5))
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if not series.get("snr", {}).get("sample_count"):
        document.add_paragraph("SNR 曲线不可用：SDB 中未发现 Product Service SNR 或带 SNR 语义的 Debug 通道。")
    if not any(series.get(role, {}).get("sample_count") for role in ("roll", "pitch", "yaw")):
        document.add_paragraph("姿态曲线不可用：SDB 中未发现可识别的 roll、pitch、yaw 数据。")
    if not analysis.get("ms6222_comparison_available"):
        document.add_paragraph(
            "MS6222 同期比较数据不可用，因此本报告不作姿态精度、收敛或发散判定。"
        )

    document.add_heading("四类工况分析", level=2)
    for scenario in analysis.get("scenario_groups", []):
        document.add_heading(
            str(scenario.get("label") or scenario.get("scenario_id")),
            level=3,
        )
        state_text = {
            "not_enabled": "未启用",
            "time_window_missing": "缺少时间窗",
            "curve_data_missing": "缺少曲线数据",
            "analyzed": "已分析",
        }.get(str(scenario.get("state")), str(scenario.get("state", "未记录")))
        _add_key_value_table(
            document,
            (
                ("工况状态", state_text),
                ("自动结果", scenario.get("automatic_result", "未记录")),
                ("测试项", "、".join(scenario.get("test_ids", []))),
                ("曲线窗口", f"{scenario.get('start_s', '-')} ～ {scenario.get('end_s', '-')} s"),
            ),
        )
        base = f"analysis/scenarios/{scenario['scenario_id']}"
        for suffix, caption in (
            ("snr.png", "SNR 与锁定状态"),
            ("attitude.png", "设备三轴姿态与 MS6222 姿态误差"),
        ):
            payload = charts.get(f"{base}/{suffix}")
            if not payload:
                continue
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(BytesIO(payload), width=Cm(16.5))
            caption_paragraph = document.add_paragraph(
                f"{scenario['label']} · {caption}"
            )
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for stage in scenario.get("stages", []):
            document.add_heading(
                f"{stage['test_id']} · 第 {stage['ordinal']} 次 · {str(stage['status']).upper()}",
                level=4,
            )
            _add_stage_statistics(document, stage, series)
        conclusion = document.add_paragraph()
        conclusion.add_run("工况结论：").bold = True
        conclusion.add_run(str(scenario.get("conclusion") or "未生成结论。"))

    document.add_heading("跨工况指标汇总", level=2)
    cross_rows = [("测试项目", "次数", "结果", "说明")]
    for stage in analysis.get("cross_scenario_stages", []):
        cross_rows.append(
            (
                stage["test_id"],
                stage["ordinal"],
                str(stage["status"]).upper(),
                "引用相应工况时间窗，不作为独立运行工况",
            )
        )
    if len(cross_rows) == 1:
        cross_rows.append(("未启用", "-", "N/A", "本批次没有跨工况导航汇总项"))
    _add_table(document, cross_rows)


def _add_stage_statistics(
    document: Document,
    stage: Mapping[str, Any],
    series: Mapping[str, Any],
) -> None:
    rows = [("曲线", "样本数", "均值", "标准差", "起始中位数", "末尾中位数", "变化", "趋势")]
    for role in ("snr", "roll", "pitch", "yaw", "attitude_error"):
        stats = stage.get("statistics", {}).get(role, {})
        if not stats.get("available"):
            continue
        rows.append(
            (
                series.get(role, {}).get("label", role),
                stats["sample_count"],
                f"{stats['mean']:.3f}",
                f"{stats['standard_deviation']:.3f}",
                f"{stats['start_median']:.3f}",
                f"{stats['end_median']:.3f}",
                f"{stats['delta']:+.3f}",
                {"stable": "稳定", "increasing": "上升", "decreasing": "下降"}.get(stats["trend"], stats["trend"]),
            )
        )
    if len(rows) == 1:
        rows.append(("未记录", "0", "-", "-", "-", "-", "-", "不可分析"))
    _add_table(document, rows)
    conclusion = document.add_paragraph()
    conclusion.add_run("子场景结论：").bold = True
    conclusion.add_run(str(stage.get("conclusion") or "未生成结论。"))


def _add_results(document: Document, snapshot: Mapping[str, Any]) -> None:
    document.add_heading("四、自动测试结果", level=1)
    attempts = snapshot["attempts"]
    if not attempts:
        document.add_paragraph("未记录测试 attempt。")
        return
    for attempt in attempts:
        document.add_heading(
            f"{attempt['test_id']}  ·  第 {attempt['ordinal']} 次  ·  {str(attempt['status']).upper()}",
            level=2,
        )
        _add_key_value_table(
            document,
            (
                ("阶段", attempt.get("phase") or "未记录"),
                ("开始单调时间", attempt.get("started_monotonic_ns") or "未记录"),
                ("结束单调时间", attempt.get("ended_monotonic_ns") or "未记录"),
                ("共享夹具动作", attempt.get("fixture_action_id") or "无"),
                ("结果详情", _compact_json(attempt.get("result_json", {}))),
            ),
        )
        metric_rows = [("指标", "值", "单位", "判定", "样本数", "证据等级")]
        for metric in attempt.get("metrics", []):
            metric_rows.append(
                (
                    metric["name"],
                    _compact_json(metric["value_json"]),
                    metric.get("unit", ""),
                    metric.get("verdict", ""),
                    metric.get("sample_count") if metric.get("sample_count") is not None else "-",
                    metric.get("evidence_level", ""),
                )
            )
        if len(metric_rows) == 1:
            metric_rows.append(("未记录", "-", "-", "-", "-", "-"))
        _add_table(document, metric_rows)
        document.add_paragraph("证据索引：device/metrics.json、device/events.jsonl、test/attempts.json")


def _add_integrity(document: Document, snapshot: Mapping[str, Any]) -> None:
    document.add_heading("五、数据完整性", level=1)
    device = snapshot["device"]
    attempts = snapshot["attempts"]
    events = snapshot["events"]
    incomplete = [item for item in attempts if item["status"] in {"incomplete", "aborted"}]
    sdb_present = bool(
        list(
            (
                Path(snapshot["batch"]["output_dir"])
                / "devices"
                / str(device["serial_number"])
                / "evidence"
            ).glob("*.sdb")
        )
    )
    disconnects = sum(
        1 for event in events if "disconnect" in str(event.get("event_type", "")).lower()
    )
    _add_key_value_table(
        document,
        (
            ("身份确认", "已记录" if device.get("device_uid") or device.get("serial_number") else "缺失"),
            ("原始 SDB", "已归档" if sdb_present else "未记录"),
            ("断线相关事件", disconnects),
            ("事件总数", len(events)),
            ("不完整测试项", len(incomplete)),
            ("时间对齐", "使用主机 UTC 与单调时钟"),
            ("完整性结论", "完整" if sdb_present and not incomplete else "存在数据缺口"),
        ),
    )
    document.add_paragraph(
        "共享夹具证据按本设备测试窗口归档；当前批次未提供的证据文件在 ZIP 中以 status=unavailable 明确标识。"
    )


def _add_review(
    document: Document,
    snapshot: Mapping[str, Any],
    version: int,
    review: Optional[Mapping[str, Any]],
) -> None:
    document.add_heading("六、审核与版本记录", level=1)
    rows = [("版本", "结论", "生成/审核时间", "审核人", "原因")]
    rows.append(("V01", snapshot["automatic_verdict"], snapshot["frozen_utc"], "自动生成", "自动测试结论"))
    if review is not None:
        rows.append(
            (
                "V02",
                review["final_verdict"],
                review["reviewed_utc"],
                review["reviewer"],
                review["reason"],
            )
        )
    else:
        rows.append(("V02", "未生成", "-", "未审核", "-"))
    _add_table(document, rows)
    document.add_paragraph(
        f"当前文档版本：V{version:02d}。自动结论始终保留，审核仅形成新的正式版本。"
    )


def _add_appendix(document: Document, snapshot: Mapping[str, Any]) -> None:
    document.add_heading("技术附录", level=1)
    document.add_heading("A. 完整指标", level=2)
    metric_rows = [("测试项目", "次数", "指标", "值", "单位", "判定", "算法版本")]
    for attempt in snapshot["attempts"]:
        for metric in attempt.get("metrics", []):
            metric_rows.append(
                (
                    attempt["test_id"],
                    attempt["ordinal"],
                    metric["name"],
                    _compact_json(metric["value_json"]),
                    metric.get("unit", ""),
                    metric.get("verdict", ""),
                    metric.get("algorithm_version", ""),
                )
            )
    if len(metric_rows) == 1:
        metric_rows.append(("未记录", "-", "-", "-", "-", "-", "-"))
    _add_table(document, metric_rows)
    document.add_heading("B. 异常事件索引", level=2)
    event_rows = [("ID", "UTC", "类型", "详情")]
    exception_events = _report_exception_events(snapshot["events"])
    for event in exception_events:
        event_rows.append(
            (
                event["event_id"],
                event["wall_time_utc"],
                event["event_type"],
                _compact_json(event["payload_json"]),
            )
        )
    if len(event_rows) == 1:
        event_rows.append(("-", "-", "无异常事件", "完整事件流见 device/events.jsonl 和 test/timeline.json"))
    _add_table(document, event_rows)
    document.add_paragraph(
        f"批次共记录 {len(snapshot['events'])} 条事件，报告列出 {len(exception_events)} 条异常相关事件。"
    )
    document.add_heading("C. 证据包索引", level=2)
    for path in (
        "device/evidence.sdb",
        "device/identity.json",
        "device/parameters.json",
        "device/metrics.json",
        "device/events.jsonl",
        "fixture/power_config.json",
        "fixture/power_identity.json",
        "fixture/power_scpi.jsonl",
        "fixture/motion_commands.jsonl",
        "fixture/ms6222_raw.jsonl",
        "fixture/ms6222_parsed.jsonl",
        "fixture/attitude_comparison.csv",
        "fixture/calibration.json",
        "test/recipe.json",
        "test/attempts.json",
        "test/timeline.json",
        "test/software.json",
        "analysis/analysis.json",
        "analysis/snr_timeline.png",
        "analysis/attitude_timeline.png",
        "analysis/convergence_timeline.png",
    ):
        paragraph = document.add_paragraph(path, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
    scenario_paths = [
        f"analysis/scenarios/{scenario['scenario_id']}/{{snr,attitude}}.png"
        for scenario in snapshot.get("analysis", {}).get("scenario_groups", [])
    ]
    if scenario_paths:
        paragraph = document.add_paragraph(
            "工况曲线（按可用数据生成）：" + "；".join(scenario_paths),
            style="List Bullet",
        )
        paragraph.paragraph_format.space_after = Pt(0)


def _set_paragraph_bottom_border(paragraph: Any) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "7F7F7F")
    borders.append(bottom)


def _add_signature_table(
    document: Document,
    values: Sequence[tuple[object, object]],
) -> None:
    table = document.add_table(rows=len(values), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(3.2)
    table.columns[1].width = Cm(9.5)
    for row_index, (label, value) in enumerate(values):
        row = table.rows[row_index]
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        for column_index, content in enumerate((label, value)):
            cell = row.cells[column_index]
            cell.width = Cm(3.2 if column_index == 0 else 9.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(str(content))
            run.font.size = Pt(10)
            if column_index == 0:
                run.bold = True


def _cover_date(batch: Mapping[str, Any]) -> str:
    value = batch.get("finished_utc") or batch.get("created_utc")
    local_date = _parse_datetime(value).astimezone()
    return f"{local_date.year} 年 {local_date.month} 月 {local_date.day} 日"


def _network_endpoint(host: object, port: object) -> str:
    host_text = str(host or "").strip()
    port_text = str(port or "").strip()
    if not host_text:
        return "未配置"
    return f"{host_text}:{port_text}" if port_text else host_text


def _environment_value(value: object, unit: str) -> str:
    if value is None or str(value).strip() == "":
        return "未记录"
    return f"{value} {unit}"


def _format_utc_offset(value: object) -> str:
    text = str(value or "").strip()
    if re.fullmatch(r"[+-]\d{4}", text):
        return f"{text[:3]}:{text[3:]}"
    return text or "未配置"


def _add_key_value_table(document: Document, values: Sequence[tuple[object, object]]) -> None:
    rows = [("项目", "内容"), *((str(key), str(value)) for key, value in values)]
    _add_table(document, rows, widths=(4.2, 12.2))


def _add_table(
    document: Document,
    rows: Sequence[Sequence[object]],
    *,
    widths: Sequence[float] = (),
) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = not bool(widths)
    for row_index, row in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            row_properties.append(repeat_header)
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            if widths and column_index < len(widths):
                cell.width = Cm(widths[column_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shading)


def _parse_datetime(value: object) -> datetime:
    text = str(value or "")
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return datetime.now(timezone.utc)
    return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)


def _time_range(batch: Mapping[str, Any]) -> str:
    return f"{batch.get('started_utc') or '未记录'} ～ {batch.get('finished_utc') or '未记录'}"


def _pretty_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n"


def _json_line(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n"


def _compact_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _read_json_file(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _read_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except (OSError, UnicodeError):
        return ""


def _first_text(root: Path, names: Sequence[str]) -> str:
    for name in names:
        direct = root / name
        if direct.is_file():
            return _read_text_file(direct)
        matches = tuple(root.glob(f"**/{name}"))
        if matches:
            return _read_text_file(matches[0])
    return _json_line({"status": "unavailable"})


def _first_text_or_csv_jsonl(
    root: Path,
    text_names: Sequence[str],
    csv_names: Sequence[str],
) -> str:
    text = _first_text(root, text_names)
    if text != _json_line({"status": "unavailable"}):
        return text
    for name in csv_names:
        matches = (root / name, *sorted(root.glob(f"**/{name}")))
        for path in matches:
            if not path.is_file():
                continue
            try:
                with path.open("r", encoding="utf-8", newline="") as stream:
                    rows = list(csv.DictReader(stream))
            except (OSError, UnicodeError, csv.Error):
                continue
            if rows:
                return "".join(_json_line(row) for row in rows)
    return text


def _first_json(root: Path, names: Sequence[str]) -> dict[str, Any]:
    for name in names:
        direct = root / name
        if direct.is_file():
            return _read_json_file(direct)
        matches = tuple(root.glob(f"**/{name}"))
        if matches:
            return _read_json_file(matches[0])
    return {"status": "unavailable"}


def _fixture_lines(actions: Sequence[Mapping[str, Any]], keyword: str) -> str:
    selected = [
        item for item in actions if keyword in str(item.get("action_type", "")).lower()
    ]
    if not selected:
        return _json_line({"status": "unavailable"})
    return "".join(_json_line(item) for item in selected)


def _report_exception_events(
    events: Sequence[Mapping[str, Any]],
) -> list[Mapping[str, Any]]:
    markers = (
        "fail",
        "error",
        "disconnect",
        "gap",
        "abort",
        "warning",
        "timeout",
        "conflict",
        "missing",
        "protect",
        "incomplete",
    )
    selected = []
    for event in events:
        event_type = str(event.get("event_type", "")).lower()
        payload = event.get("payload_json", {})
        payload_values = (
            " ".join(str(value) for value in payload.values()).lower()
            if isinstance(payload, Mapping)
            else str(payload).lower()
        )
        has_explicit_failure = any(
            value in payload_values.split()
            for value in ("fail", "failed", "error", "aborted", "incomplete")
        )
        if any(marker in event_type for marker in markers) or has_explicit_failure:
            selected.append(event)
    return selected


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _zip_entry_hashes(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    with zipfile.ZipFile(path, "r") as archive:
        for name in archive.namelist():
            result[name] = hashlib.sha256(archive.read(name)).hexdigest()
    return result


def _atomic_write_text(path: Path, value: str) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(value, encoding="utf-8")
    temporary.replace(path)


__all__ = [
    "ProductionReportService",
    "ReportArtifact",
    "ReportBranding",
    "automatic_device_verdict",
    "sanitize_path_component",
]
