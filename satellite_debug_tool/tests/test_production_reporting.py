"""Per-device DOCX report, evidence archive, review, and manifest regression tests."""

from __future__ import annotations

import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import struct
import zipfile

from docx import Document
import pytest

from satellite_debug_tool.core.production import (
    AttemptStatus,
    BatchStatus,
    ProductionRecipe,
    ProductionReportService,
    ProductionResultStore,
    ReportBranding,
    ResultStoreError,
    sanitize_path_component,
)
from satellite_debug_tool.core.production.report_analysis import (
    build_report_analysis,
    render_report_charts,
)
from satellite_debug_tool.core.protocol import CmdType, build_frame
from satellite_debug_tool.io.data_recorder import DataRecorder, SDB_VERSION_V3
from satellite_debug_tool.tests.test_production_recipe import valid_recipe


def _finish_attempt(
    store: ProductionResultStore,
    attempt_id: int,
    verdict: AttemptStatus,
) -> None:
    for status in (
        AttemptStatus.WAITING_PREREQUISITE,
        AttemptStatus.ARMED,
        AttemptStatus.RUNNING,
        AttemptStatus.ANALYZING,
        verdict,
    ):
        store.transition_attempt(
            attempt_id,
            status,
            result={"reason": "sample result"} if status == verdict else None,
        )


def _terminal_batch(
    tmp_path: Path,
    verdicts: tuple[AttemptStatus, ...],
    *,
    test_ids: tuple[str, ...] | None = None,
) -> tuple[ProductionResultStore, Path, str]:
    output = tmp_path / "raw" / "BATCH-01"
    output.mkdir(parents=True)
    recipe = ProductionRecipe.from_mapping(valid_recipe())
    store = ProductionResultStore(output / "batch.sqlite3")
    store.create_batch(
        "BATCH-01",
        recipe,
        operator="测试员甲",
        output_dir=output,
    )
    serial = "AFD01C-20260919-001"
    store.register_device(
        "BATCH-01",
        serial_number=serial,
        slot=1,
        endpoint_ip="192.168.1.12",
        endpoint_port=4004,
        hardware_type="AFD01C",
        device_uid="111111112222222233333333",
        mac_address="02:00:00:00:00:01",
        metadata={"main_firmware": "1.2.3", "parameters": {"mode": 1}},
    )
    store.start_batch(
        "BATCH-01",
        (serial,),
        test_ids or tuple(f"test-{index}" for index in range(len(verdicts))),
    )
    attempts = store.list_attempts("BATCH-01")
    for attempt, verdict in zip(attempts, verdicts):
        _finish_attempt(store, int(attempt["attempt_id"]), verdict)
        store.record_metric(
            int(attempt["attempt_id"]),
            "snr_median",
            12.5,
            unit="dB",
            verdict=verdict.value,
            sample_count=100,
        )
    target_status = (
        BatchStatus.INCOMPLETE
        if any(value in {AttemptStatus.INCOMPLETE, AttemptStatus.ABORTED} for value in verdicts)
        else BatchStatus.COMPLETED
    )
    store.transition_batch("BATCH-01", target_status)
    evidence_dir = output / "devices" / serial / "evidence"
    evidence_dir.mkdir(parents=True)
    (evidence_dir / "capture.sdb").write_bytes(b"SDB3-example")
    (output / "power_supply_records.jsonl").write_text(
        '{"command":"MEAS:VOLT?","response":"12.01"}\n', encoding="utf-8"
    )
    (output / "power_supply_summary.json").write_text(
        json.dumps({"identity": {"model": "PSW80-27"}}), encoding="utf-8"
    )
    (output / "production_configuration.json").write_text(
        json.dumps(
            {
                "product_template": {
                    "display_name": "AFD01C 卫星通信相控阵终端",
                    "product": "afd01c",
                },
                "power_profile": {"manufacturer": "GW-INSTEK", "model": "PSW80-27"},
                "station_profile": {
                    "display_name": "一号工位",
                    "motion_profile_id": "motion-main",
                    "reference_profile_id": "ms6222-main",
                },
                "combined_supply": {"voltage_v": 12.0, "current_a": 8.0},
                "test_environment": {
                    "location": "生产测试实验室",
                    "temperature_c": 23.5,
                    "relative_humidity_percent": 48,
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return store, output, serial


def _document_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _service_fast(timestamp_ms: int, snr_db: float, roll_deg: float) -> bytes:
    payload = struct.pack(
        "<BII6B6f",
        1,
        timestamp_ms,
        0xFFF,
        0,
        3,
        1,
        3,
        3,
        0,
        roll_deg,
        2.0,
        3.0,
        120.0,
        35.0,
        snr_db,
    )
    return build_frame(CmdType.SERVICE_FAST_STATE, payload)


def test_curve_analysis_uses_sdb_host_time_and_generates_stage_charts(tmp_path: Path) -> None:
    path = tmp_path / "capture.sdb"
    recorder = DataRecorder(path, format_version=SDB_VERSION_V3)
    assert recorder.start()
    wall_origin = 1_700_000_000_000_000_000
    for index, snr in enumerate((8.0, 10.0, 12.0, 14.0, 16.0)):
        assert recorder.write_frame(
            _service_fast(index * 1000, snr, float(index)),
            host_timestamp_ns=wall_origin + index * 1_000_000_000,
        )
    assert recorder.stop()

    monotonic_origin = 5_000_000_000
    events = [
        {
            "monotonic_ns": monotonic_origin,
            "wall_time_utc": "2023-11-14T22:13:20+00:00",
        }
    ]
    attempts = [
        {
            "attempt_id": 1,
            "test_id": "locked_rocking",
            "ordinal": 1,
            "status": "pass",
            "phase": "observing",
            "started_monotonic_ns": monotonic_origin,
            "ended_monotonic_ns": monotonic_origin + 4_000_000_000,
        }
    ]
    analysis = build_report_analysis((path,), attempts, events, tmp_path)
    stats = analysis["stages"][0]["statistics"]["snr"]

    assert analysis["source_kind"] == "product_service"
    assert stats["sample_count"] == 5
    assert stats["trend"] == "increasing"
    assert stats["start_median"] == pytest.approx(8.0)
    assert stats["end_median"] == pytest.approx(16.0)
    assert analysis["ms6222_comparison_available"] is False
    charts = render_report_charts(analysis)
    assert charts["analysis/snr_timeline.png"].startswith(b"\x89PNG")
    assert charts["analysis/attitude_timeline.png"].startswith(b"\x89PNG")


def test_report_groups_five_test_ids_into_four_operating_scenarios(tmp_path: Path) -> None:
    scenario_ids = (
        "static_acquisition",
        "locked_rocking",
        "power_on_rocking",
        "locked_drive",
        "power_on_drive",
        "gnss",
    )
    attempts = [
        {
            "attempt_id": index,
            "test_id": test_id,
            "ordinal": 1,
            "status": "pass",
            "phase": "observing",
            "started_monotonic_ns": index * 1_000_000_000,
            "ended_monotonic_ns": (index + 1) * 1_000_000_000,
        }
        for index, test_id in enumerate(scenario_ids, start=1)
    ]
    events = [{"monotonic_ns": 0, "wall_time_utc": "2023-11-14T22:13:20+00:00"}]
    analysis = build_report_analysis((), attempts, events, tmp_path)

    assert [item["scenario_id"] for item in analysis["scenario_groups"]] == [
        "static",
        "locked_rocking",
        "power_on_rocking",
        "vehicle_dynamic",
    ]
    vehicle = analysis["scenario_groups"][-1]
    assert [item["test_id"] for item in vehicle["stages"]] == [
        "locked_drive",
        "power_on_drive",
    ]
    assert [item["test_id"] for item in analysis["cross_scenario_stages"]] == ["gnss"]


def test_valid_process_recording_is_embedded_in_docx_and_evidence(tmp_path: Path) -> None:
    store, output, serial = _terminal_batch(
        tmp_path,
        (AttemptStatus.PASS,),
        test_ids=("static_acquisition",),
    )
    try:
        path = output / "devices" / serial / "evidence" / "capture.sdb"
        path.unlink()
        stored_attempt = store.list_attempts("BATCH-01")[0]
        stored_events = store.list_events("BATCH-01")
        offsets = [
            int(datetime.fromisoformat(event["wall_time_utc"]).timestamp() * 1_000_000_000)
            - int(event["monotonic_ns"])
            for event in stored_events
        ]
        wall_offset = sorted(offsets)[len(offsets) // 2]
        attempt_start = int(stored_attempt["started_monotonic_ns"])
        attempt_end = int(stored_attempt["ended_monotonic_ns"])
        recorder = DataRecorder(path, format_version=SDB_VERSION_V3)
        assert recorder.start()
        for index in range(6):
            fraction = index / 5
            sample_monotonic = int(
                attempt_start + (attempt_end - attempt_start) * fraction
            )
            assert recorder.write_frame(
                _service_fast(index * 1000, 10.0 + index, float(index)),
                host_timestamp_ns=sample_monotonic + wall_offset,
            )
        assert recorder.stop()
        fixture_dir = output / "fixture_capture"
        fixture_dir.mkdir()
        comparison_fields = (
            "monotonic_ns",
            "target_roll_deg",
            "target_pitch_deg",
            "target_yaw_deg",
            "measured_roll_deg",
            "measured_pitch_deg",
            "measured_yaw_deg",
            "error_roll_deg",
            "error_pitch_deg",
            "error_yaw_deg",
            "error_angle_deg",
        )
        event_origin = attempt_start
        with (fixture_dir / "attitude_comparison.csv").open(
            "w", encoding="utf-8", newline=""
        ) as stream:
            writer = csv.DictWriter(stream, fieldnames=comparison_fields)
            writer.writeheader()
            for index, error in enumerate((4.0, 3.0, 2.0, 1.0)):
                writer.writerow(
                    {
                        "monotonic_ns": event_origin + index * 1_000_000_000,
                        "target_roll_deg": 0,
                        "target_pitch_deg": 0,
                        "target_yaw_deg": 0,
                        "measured_roll_deg": error,
                        "measured_pitch_deg": 0,
                        "measured_yaw_deg": 0,
                        "error_roll_deg": error,
                        "error_pitch_deg": 0,
                        "error_yaw_deg": 0,
                        "error_angle_deg": error,
                    }
                )
        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
        )
        artifact = service.archive_device("BATCH-01", serial)

        assert artifact.document_path is not None
        document = Document(artifact.document_path)
        assert len(document.inline_shapes) >= 3
        assert "三、过程曲线与阶段分析" in _document_text(artifact.document_path)
        assert "静态工况" in _document_text(artifact.document_path)
        assert "车载动态工况" in _document_text(artifact.document_path)
        with zipfile.ZipFile(artifact.evidence_path) as archive:
            names = set(archive.namelist())
            assert "analysis/analysis.json" in names
            assert "analysis/snr_timeline.png" in names
            assert "analysis/attitude_timeline.png" in names
            assert "analysis/convergence_timeline.png" in names
            assert "analysis/scenarios/static/snr.png" in names
            assert "analysis/scenarios/static/attitude.png" in names
    finally:
        store.close()


def test_pass_generates_v1_self_contained_zip_and_valid_manifest(tmp_path: Path) -> None:
    store, output, serial = _terminal_batch(tmp_path, (AttemptStatus.PASS,))
    try:
        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
            branding=ReportBranding(company_name="软赫", tester_role="测试员"),
        )
        artifact = service.archive_device("BATCH-01", serial)

        assert artifact.status == "generated_v1"
        assert artifact.document_path is not None and artifact.document_path.is_file()
        assert artifact.directory.parts[-3:] == ("AFD01C", serial, "BATCH-01")
        with zipfile.ZipFile(artifact.evidence_path) as archive:
            names = set(archive.namelist())
            assert {
                "device/evidence.sdb",
                "device/identity.json",
                "device/metrics.json",
                "fixture/power_scpi.jsonl",
                "fixture/ms6222_raw.jsonl",
                "test/recipe.json",
                "test/software.json",
            } <= names
            assert archive.read("device/evidence.sdb") == b"SDB3-example"

        manifest = json.loads(artifact.manifest_path.read_text(encoding="utf-8"))
        assert manifest["automatic_verdict"] == "PASS"
        assert manifest["evidence"]["sha256"] == hashlib.sha256(
            artifact.evidence_path.read_bytes()
        ).hexdigest()
        assert manifest["reports"][0]["sha256"] == hashlib.sha256(
            artifact.document_path.read_bytes()
        ).hexdigest()
        text = _document_text(artifact.document_path)
        assert "测试报告" in text
        assert "卫星通信相控阵设备自动化测试" in text
        assert "产品名称" in text
        assert "AFD01C 卫星通信相控阵终端" in text
        assert "编写" in text and "审核" in text and "批准" in text
        assert "测试环境" in text
        assert "生产测试实验室" in text
        assert "环境温度" in text and "23.5 °C" in text
        assert "相对湿度" in text and "48 %RH" in text
        assert "操作系统" in text
        assert "一、测试摘要" in text
        assert "技术附录" in text
        assert serial in text
    finally:
        store.close()


def test_fail_v1_can_create_one_audited_v2_without_overwriting_v1(tmp_path: Path) -> None:
    store, output, serial = _terminal_batch(tmp_path, (AttemptStatus.FAIL,))
    try:
        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
        )
        first = service.archive_device("BATCH-01", serial)
        first_bytes = first.document_path.read_bytes() if first.document_path else b""
        second = service.review(
            "BATCH-01",
            serial,
            reviewer="审核员乙",
            final_verdict="PASS",
            reason="完成线下复测并确认原始阈值配置错误。",
        )

        assert second.document_path is not None and second.document_path.is_file()
        assert first.document_path is not None and first.document_path.read_bytes() == first_bytes
        text = _document_text(second.document_path)
        assert "审核员乙" in text
        assert "完成线下复测" in text
        assert "FAIL" in text and "PASS" in text
        with pytest.raises(ResultStoreError, match="already exists"):
            service.review(
                "BATCH-01",
                serial,
                reviewer="另一审核员",
                final_verdict="FAIL",
                reason="重复审核",
            )
    finally:
        store.close()


def test_incomplete_only_archives_evidence_and_cannot_be_reviewed(tmp_path: Path) -> None:
    store, output, serial = _terminal_batch(tmp_path, (AttemptStatus.INCOMPLETE,))
    try:
        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
        )
        artifact = service.archive_device("BATCH-01", serial)
        assert artifact.status == "not_generated"
        assert artifact.document_path is None
        assert artifact.evidence_path.is_file()
        assert artifact.manifest_path.is_file()
        assert not tuple(artifact.directory.glob("*.docx"))
        with pytest.raises(ResultStoreError, match="cannot be approved"):
            store.create_report_review(
                "BATCH-01",
                serial,
                reviewer="审核员",
                final_verdict="PASS",
                reason="不得通过",
            )
    finally:
        store.close()


def test_path_components_are_safe_and_bounded() -> None:
    assert sanitize_path_component(' AFD01C:/<>"\\|?* ') == "AFD01C"
    assert sanitize_path_component("..") == "unknown"
    assert len(sanitize_path_component("A" * 500)) == 128


def test_two_devices_receive_separate_self_contained_archives(tmp_path: Path) -> None:
    output = tmp_path / "raw" / "BATCH-02"
    output.mkdir(parents=True)
    recipe = ProductionRecipe.from_mapping(valid_recipe())
    store = ProductionResultStore(output / "batch.sqlite3")
    serials = ("AFD01C-A", "AFD01C-B")
    try:
        store.create_batch(
            "BATCH-02", recipe, operator="operator", output_dir=output
        )
        for slot, serial in enumerate(serials, start=1):
            store.register_device(
                "BATCH-02",
                serial_number=serial,
                slot=slot,
                hardware_type="AFD01C",
            )
        store.start_batch("BATCH-02", serials, ("identity",))
        for attempt in store.list_attempts("BATCH-02"):
            _finish_attempt(store, int(attempt["attempt_id"]), AttemptStatus.PASS)
        store.transition_batch("BATCH-02", BatchStatus.COMPLETED)
        for serial in serials:
            evidence = output / "devices" / serial / "evidence"
            evidence.mkdir(parents=True)
            (evidence / "capture.sdb").write_bytes(serial.encode("ascii"))

        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
        )
        artifacts = service.archive_batch("BATCH-02")
        assert [item.serial_number for item in artifacts] == list(serials)
        assert artifacts[0].directory != artifacts[1].directory
        for artifact in artifacts:
            with zipfile.ZipFile(artifact.evidence_path) as archive:
                assert archive.read("device/evidence.sdb") == artifact.serial_number.encode(
                    "ascii"
                )
            assert artifact.document_path is not None
            report_text = _document_text(artifact.document_path)
            assert "测试环境" in report_text
            assert "测试地点\n未配置" in report_text
            assert "环境温度\n未记录" in report_text
    finally:
        store.close()


def test_docx_failure_records_generation_failure_without_changing_verdict(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    store, output, serial = _terminal_batch(tmp_path, (AttemptStatus.PASS,))
    try:
        service = ProductionReportService(
            store,
            batch_output_dir=output,
            report_root=tmp_path / "reports",
        )

        def fail_document(*_args, **_kwargs):
            raise OSError("simulated docx failure")

        monkeypatch.setattr(service, "_write_docx", fail_document)
        with pytest.raises(OSError, match="simulated"):
            service.archive_device("BATCH-01", serial)

        assert store.get_batch("BATCH-01")["status"] == BatchStatus.COMPLETED.value
        snapshot = store.get_report_snapshot("BATCH-01", serial)
        assert snapshot["automatic_verdict"] == "PASS"
        artifact = store.get_report_artifact("BATCH-01", serial, 1)
        assert artifact["status"] == "generation_failed"
        assert "simulated docx failure" in artifact["error"]
    finally:
        store.close()
